# -*- encoding: utf-8 -*-
##############################################################################
#
# ERP Heritage
# Copyright (C) 2026 (https://www.erpheritage.com.au/)
#
##############################################################################
"""
Multi-sheet XLSX export of the daily disbursement dashboard.

eh_account_base's XlsxReportWriter renders ONE sheet in the dynamic-report
payload shape; this export needs five sheets with different column sets,
so it is its own small writer that follows the same conventions (Arial,
grey header fill, accounting number format with red negatives).

Plain Python: takes already-computed figures and returns bytes, so it is
unit-testable without the ORM. All labels are passed in (translated by
the caller).

Sheets: Summary, By Day, Products, By Tag, Lines.
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONEY = '#,##0.00;[Red]-#,##0.00'
QTY = '#,##0.##'
PERCENT = '+0.0%;[Red]-0.0%;0.0%'
DATE = 'yyyy-mm-dd'


class DisbursementWorkbook:

    BOLD = Font(name='Arial', size=10, bold=True)
    TITLE = Font(name='Arial', size=14, bold=True)
    META = Font(name='Arial', size=9, italic=True, color='666666')
    HEADER_FILL = PatternFill(start_color='F0F0F0', end_color='F0F0F0', fill_type='solid')
    TOP_BORDER = Border(top=Side(border_style='thin', color='000000'))

    def __init__(self, title, labels):
        self.title = title
        self.l = labels
        self.workbook = Workbook()

    def build(self, data, filters, days, products, lines):
        self._summary(self.workbook.active, data, filters)
        self._table(self.l['by_day'],
                    [(self.l['date'], DATE), (self.l['disbursed'], MONEY),
                     (self.l['returned'], MONEY), (self.l['net'], MONEY)],
                    [(date.fromisoformat(p['date']), p['disbursed'], p['returned'], p['net'])
                     for p in days],
                    totals=True)
        self._table(self.l['products'],
                    [(self.l['product'], None), (self.l['uom'], None),
                     (self.l['out_qty'], QTY), (self.l['out_cost'], MONEY),
                     (self.l['return_qty'], QTY), (self.l['return_cost'], MONEY),
                     (self.l['net_qty'], QTY), (self.l['net_cost'], MONEY)],
                    [(p['name'], p['uom'], p['out_qty'], p['out_cost'], p['return_qty'],
                      p['return_cost'], p['net_qty'], p['net_cost']) for p in products],
                    totals=True)
        self._table(self.l['by_tag'],
                    [(self.l['tag'], None), (self.l['disbursed'], MONEY),
                     (self.l['returned'], MONEY), (self.l['net'], MONEY)],
                    [(t['name'], t['disbursed_cost'], t['returned_cost'], t['net_cost'])
                     for t in data['by_tag']],
                    totals=True)
        self._table(self.l['lines'],
                    [(self.l['date'], DATE), (self.l['sheet'], None), (self.l['state'], None),
                     (self.l['tag'], None), (self.l['product'], None), (self.l['type'], None),
                     (self.l['qty'], QTY), (self.l['uom'], None), (self.l['unit_cost'], MONEY),
                     (self.l['subtotal'], MONEY), (self.l['signed_subtotal'], MONEY),
                     (self.l['currency'], None), (self.l['company'], None), (self.l['note'], None)],
                    [(r['date'], r['sheet'], r['state'], r['tag'], r['product'], r['move_type'],
                      r['qty'], r['uom'], r['unit_cost'], r['subtotal'], r['signed_subtotal'],
                      r['currency'], r['company'], r['note']) for r in lines])
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------

    def _summary(self, ws, data, filters):
        ws.title = self.l['summary'][:31]
        period = data['period']
        ws.append([self.title])
        ws['A1'].font = self.TITLE
        meta = [
            (self.l['period'], '%s → %s' % (period['date_from'], period['date_to'])),
            (self.l['compared_with'], '%s → %s' % (period['prev_from'], period['prev_to'])),
            (self.l['tag'], filters['tag']),
            (self.l['drafts'], self.l['yes'] if filters['include_draft'] else self.l['no']),
            (self.l['company'], filters['company']),
            (self.l['currency'], data['currency']['symbol']),
        ]
        for label, value in meta:
            ws.append([label, value])
            ws.cell(ws.max_row, 1).font = self.META
        if data['converted']:
            ws.append([self.l['converted']])
            ws.cell(ws.max_row, 1).font = self.META
        ws.append([])
        ws.append([self.l['metric'], self.l['value'], self.l['previous'], self.l['change']])
        self._style_header(ws, ws.max_row, 4)
        for key in ('disbursed_cost', 'returned_cost', 'net_cost',
                    'disbursed_qty', 'returned_qty', 'net_qty', 'product_count', 'sheet_count'):
            kpi = data['kpi'][key]
            change = kpi['delta_pct'] / 100 if kpi['delta_pct'] is not None else None
            ws.append([self.l[key], kpi['value'], kpi['previous'], change])
            fmt = MONEY if key.endswith('_cost') else QTY
            for col, number_format in ((2, fmt), (3, fmt), (4, PERCENT)):
                ws.cell(ws.max_row, col).number_format = number_format
        self._widths(ws, [28, 22, 22, 12])

    def _table(self, title, columns, rows, totals=False):
        ws = self.workbook.create_sheet(title[:31])
        ws.append([name for name, _fmt in columns])
        self._style_header(ws, 1, len(columns))
        for row in rows:
            ws.append(list(row))
        for col, (_name, fmt) in enumerate(columns, start=1):
            if fmt:
                for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    cell.number_format = fmt
        last = ws.max_row
        if totals and rows:
            ws.append([self.l['total']] + [
                '=SUM(%s2:%s%d)' % (get_column_letter(col), get_column_letter(col), last)
                if fmt in (MONEY, QTY) else None
                for col, (_name, fmt) in enumerate(columns[1:], start=2)
            ])
            for col, (_name, fmt) in enumerate(columns, start=1):
                cell = ws.cell(ws.max_row, col)
                cell.font, cell.border = self.BOLD, self.TOP_BORDER
                if fmt in (MONEY, QTY):
                    cell.number_format = fmt
        ws.freeze_panes = 'A2'
        if rows:
            ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(columns)), last)
        self._widths(ws, [self._width(name, fmt) for name, fmt in columns])

    def _style_header(self, ws, row, count):
        for col in range(1, count + 1):
            cell = ws.cell(row, col)
            cell.font, cell.fill = self.BOLD, self.HEADER_FILL
            cell.alignment = Alignment(vertical='center')

    @staticmethod
    def _width(name, fmt):
        base = {MONEY: 16, QTY: 12, DATE: 12}.get(fmt, 24)
        return max(base, len(str(name)) + 2)

    @staticmethod
    def _widths(ws, widths):
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width

# -*- encoding: utf-8 -*-
##############################################################################
#
# ERP Heritage
# Copyright (C) 2026 (https://www.erpheritage.com.au/)
#
##############################################################################
"""
Import lines into a draft daily sheet from a CSV or XLSX file.

All-or-nothing: every row is validated first and the problems are reported
together with their spreadsheet row numbers; lines are created only when
the whole file is clean, so a half-imported day never exists.

Columns (header row required, case-insensitive, English or Arabic):
    product    internal reference, barcode, or exact product name (required)
    type       out / disbursement / return (optional, defaults to out)
    qty        quantity, > 0 (required)
    unit_cost  optional; empty means today's product cost
    note       optional
"""

import base64
import csv
import io

from odoo import _, fields, models
from odoo.exceptions import UserError


# Header aliases -> canonical column.
_COLUMNS = {
    'product': 'product', 'internal reference': 'product',
    'default_code': 'product', 'barcode': 'product', 'code': 'product',
    'المنتج': 'product', 'الصنف': 'product',
    'type': 'type', 'move_type': 'type', 'direction': 'type', 'النوع': 'type',
    'qty': 'qty', 'quantity': 'qty', 'الكمية': 'qty',
    'unit_cost': 'unit_cost', 'unit cost': 'unit_cost', 'cost': 'unit_cost',
    'التكلفة': 'unit_cost', 'سعر التكلفة': 'unit_cost',
    'note': 'note', 'notes': 'note', 'ملاحظة': 'note', 'ملاحظات': 'note',
}
_TYPES = {
    '': 'out', 'out': 'out', 'disbursement': 'out', 'disbursed': 'out',
    'صرف': 'out',
    'return': 'return', 'returned': 'return', 'in': 'return',
    'مرتجع': 'return',
}
_MAX_REPORTED_ERRORS = 20


class EhDailyDisbursementImport(models.TransientModel):
    _name = 'eh.daily.disbursement.import'
    _description = "Import Daily Disbursement Lines"

    sheet_id = fields.Many2one(
        'eh.daily.disbursement', required=True, ondelete='cascade',
    )
    file = fields.Binary(string="File", required=True, attachment=False)
    filename = fields.Char()

    def action_import(self):
        self.ensure_one()
        sheet = self.sheet_id
        sheet._check_group('eh_daily_disbursement.group_daily_disbursement_user', _('import into'))
        sheet._check_state(('draft',), _('import into'))
        rows = self._read_rows(base64.b64decode(self.file or b''))
        vals_list, errors = self._prepare_lines(rows)
        if errors:
            shown = errors[:_MAX_REPORTED_ERRORS]
            if len(errors) > len(shown):
                shown.append(_("... and %s more.", len(errors) - len(shown)))
            raise UserError(_("Nothing was imported. Fix these rows and try again:\n%s",
                              '\n'.join(shown)))
        if not vals_list:
            raise UserError(_("The file has no data rows."))
        sheet.write({'line_ids': [(0, 0, vals) for vals in vals_list]})
        sheet.message_post(body=_(
            "%(count)s lines imported from %(file)s.",
            count=len(vals_list), file=self.filename or _('a file'),
        ))
        return {'type': 'ir.actions.act_window_close'}

    # ------------------------------------------------------------------
    # Reading
    # ------------------------------------------------------------------

    def _read_rows(self, data):
        """Return [(row_number, {column: value})] for non-blank rows."""
        if data[:4] == b'\xd0\xcf\x11\xe0':
            raise UserError(_("Old .xls files are not supported. Save the file as .xlsx or .csv."))
        table = self._read_xlsx(data) if data[:2] == b'PK' else self._read_csv(data)
        if not table:
            raise UserError(_("The file is empty."))
        header = [_COLUMNS.get(str(h or '').strip().lower()) for h in table[0]]
        missing = {'product', 'qty'} - set(header)
        if missing:
            raise UserError(_(
                "The header row must contain the columns: %(cols)s. "
                "Download the template to see the expected layout.",
                cols=', '.join(sorted(missing)),
            ))
        rows = []
        for number, cells in enumerate(table[1:], start=2):
            row = {col: cell for col, cell in zip(header, cells) if col}
            if any(str(v).strip() for v in row.values() if v is not None):
                rows.append((number, row))
        return rows

    def _read_csv(self, data):
        for encoding in ('utf-8-sig', 'cp1256'):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise UserError(_("Could not read the CSV file. Save it as UTF-8."))
        try:
            dialect = csv.Sniffer().sniff(text.split('\n', 1)[0], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        return list(csv.reader(io.StringIO(text), dialect))

    def _read_xlsx(self, data):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(_("Reading .xlsx needs the openpyxl library; import a .csv instead.")) from None
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - any parse failure means "not a usable xlsx"
            raise UserError(_("Could not read the .xlsx file.")) from None
        return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]

    # ------------------------------------------------------------------
    # Validation
    # ------------------------------------------------------------------

    def _prepare_lines(self, rows):
        sheet = self.sheet_id
        allowed_categ_ids = None
        if sheet.tag_id.product_categ_ids:
            allowed_categ_ids = set(self.env['product.category'].search(
                [('id', 'child_of', sheet.tag_id.product_categ_ids.ids)]).ids)
        products = {}
        vals_list, errors = [], []
        for sequence, (number, row) in enumerate(rows, start=1):
            problems = []
            key = str(row.get('product') or '').strip()
            if key not in products:
                products[key] = self._find_product(key, sheet.company_id)
            product, product_error = products[key]
            if product_error:
                problems.append(product_error)
            elif allowed_categ_ids is not None and product.categ_id.id not in allowed_categ_ids:
                problems.append(_(
                    "%(product)s is outside the categories of tag %(tag)s",
                    product=product.display_name, tag=sheet.tag_id.display_name))

            move_type = _TYPES.get(str(row.get('type') or '').strip().lower())
            if not move_type:
                problems.append(_("unknown type %r (use out or return)", row.get('type')))

            qty = self._to_number(row.get('qty'))
            if not qty or qty <= 0:
                problems.append(_("quantity must be a number greater than zero"))

            unit_cost = self._to_number(row.get('unit_cost'))
            if unit_cost is False or (unit_cost is not None and unit_cost < 0):
                problems.append(_("unit cost must be empty or a number of zero or more"))

            if problems:
                errors.append(_("Row %(row)s: %(problems)s",
                                row=number, problems='; '.join(problems)))
                continue
            vals = {
                'sequence': sequence,
                'move_type': move_type,
                'product_id': product.id,
                'qty': qty,
                'note': str(row.get('note') or '').strip() or False,
            }
            if unit_cost is not None:
                vals['unit_cost'] = unit_cost
            vals_list.append(vals)
        return vals_list, errors

    def _find_product(self, key, company):
        """Return (product, error): internal reference, then barcode, then
        exact (case-insensitive) name, among storable/consumable products
        the sheet's company can use."""
        if not key:
            return None, _("product is empty")
        Product = self.env['product.product']
        base = [('type', '!=', 'service'),
                '|', ('company_id', '=', False), ('company_id', '=', company.id)]
        for field, op in (('default_code', '='), ('barcode', '='), ('name', '=ilike')):
            found = Product.search(base + [(field, op, key)], limit=2)
            if len(found) == 1:
                return found, None
            if len(found) > 1:
                return None, _("%(key)r matches several products by %(field)s",
                               key=key, field=field)
        return None, _("product %r not found", key)

    @staticmethod
    def _to_number(value):
        """float; None when the cell is empty; False when it is not a number.

        A lone decimal comma ("2,5") is accepted for locales that write it.
        """
        if value is None or (isinstance(value, str) and not value.strip()):
            return None
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
        text = str(value).strip()
        if ',' in text and '.' not in text:
            text = text.replace(',', '.')
        try:
            return float(text)
        except ValueError:
            return False

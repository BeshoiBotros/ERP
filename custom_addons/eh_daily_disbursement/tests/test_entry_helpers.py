# -*- encoding: utf-8 -*-
##############################################################################
#
# ERP Heritage
# Copyright (C) 2026 (https://www.erpheritage.com.au/)
#
##############################################################################
"""Copy-previous, line import wizard, analysis action and reminder cron."""

import base64
import io
from datetime import date, timedelta

from odoo import fields
from odoo.exceptions import AccessError, UserError
from odoo.tests import tagged

from .common import EhDailyDisbursementCase


@tagged('post_install', '-at_install', 'eh_daily_disbursement')
class TestCopyPrevious(EhDailyDisbursementCase):

    def test_copies_lines_with_todays_cost(self):
        previous = self._create_sheet(
            [('out', self.product_a, 3), ('return', self.product_b, 1)],
            date=date(2026, 6, 1))
        self.product_a.standard_price = 12.0
        sheet = self._create_sheet(date=date(2026, 6, 3))
        sheet.with_user(self.user).action_copy_previous()
        self.assertEqual(
            sorted(sheet.line_ids.mapped(lambda l: (l.move_type, l.product_id.id, l.qty))),
            sorted(previous.line_ids.mapped(lambda l: (l.move_type, l.product_id.id, l.qty))))
        self.assertEqual(
            sheet.line_ids.filtered(lambda l: l.product_id == self.product_a).unit_cost, 12.0)
        self.assertIn(previous.name, sheet.message_ids[0].body)

    def test_picks_latest_matching_sheet(self):
        self._create_sheet([('out', self.product_a, 1)], date=date(2026, 6, 1))
        latest = self._create_sheet([('out', self.product_a, 7)], date=date(2026, 6, 2))
        # Ignored: cancelled, other tag, later date.
        cancelled = self._create_sheet([('out', self.product_a, 99)], date=date(2026, 6, 4))
        cancelled.action_cancel()
        self._create_sheet([('out', self.product_a, 50)], date=date(2026, 6, 4), tag_id=False)
        self._create_sheet([('out', self.product_a, 60)], date=date(2026, 6, 9))
        sheet = self._create_sheet(date=date(2026, 6, 5))
        self.assertEqual(sheet._previous_sheet(), latest)
        sheet.action_copy_previous()
        self.assertEqual(sheet.line_ids.qty, 7)

    def test_refused_when_not_empty_missing_or_locked(self):
        sheet = self._create_sheet([('out', self.product_a, 1)], date=date(2026, 6, 10))
        with self.assertRaises(UserError):
            sheet.action_copy_previous()          # has lines
        lonely = self._create_sheet(date=date(2026, 6, 11), tag_id=False)
        with self.assertRaises(UserError):
            lonely.action_copy_previous()         # nothing earlier without tag
        sheet.action_confirm()
        with self.assertRaises(UserError):
            sheet.action_copy_previous()          # not draft

    def test_view_lines_action_scoped_to_sheet(self):
        sheet = self._create_sheet([('out', self.product_a, 1)], date=date(2026, 6, 12))
        action = sheet.action_view_lines()
        self.assertEqual(action['res_model'], 'eh.daily.disbursement.line')
        self.assertEqual(action['domain'], [('sheet_id', '=', sheet.id)])


@tagged('post_install', '-at_install', 'eh_daily_disbursement')
class TestImportLines(EhDailyDisbursementCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.product_a.default_code = 'PA-1'
        cls.product_b.barcode = '6221000000017'
        cls.sheet = cls._create_sheet(date=date(2026, 7, 1))

    def _import(self, content, filename='lines.csv', user=None, sheet=None):
        data = content.encode('utf-8') if isinstance(content, str) else content
        wizard = self.env['eh.daily.disbursement.import'].with_user(user or self.user).create({
            'sheet_id': (sheet or self.sheet).id,
            'file': base64.b64encode(data),
            'filename': filename,
        })
        return wizard.action_import()

    def test_csv_matches_reference_barcode_and_name(self):
        self._import(
            "product,type,qty,unit_cost,note\n"
            "PA-1,out,3,,\n"
            "6221000000017,return,2,4.5,damaged\n"
            "perfume b,,1,,\n")
        lines = self.sheet.line_ids.sorted('sequence')
        self.assertEqual(len(lines), 3)
        self.assertEqual((lines[0].product_id, lines[0].unit_cost), (self.product_a, 10.0))
        self.assertEqual((lines[1].move_type, lines[1].unit_cost, lines[1].note), ('return', 4.5, 'damaged'))
        self.assertEqual((lines[2].product_id, lines[2].move_type), (self.product_b, 'out'))
        self.assertEqual(self.sheet.net_cost, 30 - 9 + 5)
        self.assertIn('3 lines imported', self.sheet.message_ids[0].body)

    def test_semicolon_csv_with_arabic_headers_and_decimal_comma(self):
        self._import("المنتج;النوع;الكمية\nPA-1;صرف;2,5\nPA-1;مرتجع;1\n")
        self.assertEqual(self.sheet.disbursed_qty, 2.5)
        self.assertEqual(self.sheet.returned_qty, 1)

    def test_xlsx(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(['Product', 'Quantity', 'Type'])
        workbook.active.append(['PA-1', 4, 'out'])
        workbook.active.append([None, None, None])   # blank rows are skipped
        workbook.active.append(['PA-1', 1.5, 'return'])
        buffer = io.BytesIO()
        workbook.save(buffer)
        self._import(buffer.getvalue(), filename='lines.xlsx')
        self.assertEqual((self.sheet.disbursed_qty, self.sheet.returned_qty), (4, 1.5))

    def test_any_bad_row_imports_nothing(self):
        outsider = self.env['product.product'].create({
            'name': 'Hammer', 'default_code': 'HM-1', 'categ_id': self.categ_other.id})
        with self.assertRaises(UserError) as caught:
            self._import(
                "product,type,qty,unit_cost\n"
                "PA-1,out,1,\n"          # row 2 fine
                "NOPE,out,1,\n"          # row 3 unknown product
                "PA-1,sideways,1,\n"     # row 4 bad type
                "PA-1,out,0,\n"          # row 5 zero qty
                "PA-1,out,abc,\n"        # row 6 not a number
                "PA-1,out,1,-3\n"        # row 7 negative cost
                "HM-1,out,1,\n")         # row 8 outside tag categories
        message = str(caught.exception)
        for row in range(3, 9):
            self.assertIn('Row %s:' % row, message)
        self.assertNotIn('Row 2:', message)
        self.assertIn(outsider.display_name, message)
        self.assertFalse(self.sheet.line_ids)

    def test_ambiguous_name_reported(self):
        self.env['product.product'].create({'name': 'Perfume A', 'categ_id': self.categ_perfume.id})
        with self.assertRaises(UserError) as caught:
            self._import("product,qty\nPerfume A,1\n")
        self.assertIn('several products', str(caught.exception))

    def test_header_and_file_checks(self):
        with self.assertRaises(UserError):
            self._import("name,amount\nPA-1,1\n")         # missing required columns
        with self.assertRaises(UserError):
            self._import("product,qty\n")                 # no data rows
        with self.assertRaises(UserError):
            self._import(b'\xd0\xcf\x11\xe0legacy', filename='old.xls')

    def test_refused_on_locked_sheet_and_for_auditor(self):
        self._import("product,qty\nPA-1,1\n")
        self.sheet.action_confirm()
        with self.assertRaises(UserError):
            self._import("product,qty\nPA-1,1\n")
        draft = self._create_sheet(date=date(2026, 7, 2))
        with self.assertRaises(AccessError):  # no ACL on the wizard for auditors
            self._import("product,qty\nPA-1,1\n", user=self.auditor, sheet=draft)


@tagged('post_install', '-at_install', 'eh_daily_disbursement')
class TestReminderCron(EhDailyDisbursementCase):

    def setUp(self):
        super().setUp()
        self.activity_type = self.env.ref('eh_daily_disbursement.mail_activity_confirm_sheet')
        today = fields.Date.context_today(self.env['eh.daily.disbursement'])
        self.overdue = self._create_sheet(
            [('out', self.product_a, 1)], date=today - timedelta(days=1),
            responsible_id=self.user.id)
        self.current = self._create_sheet(
            [('out', self.product_a, 1)], date=today, responsible_id=self.user.id)

    def _reminders(self, sheet):
        return sheet.activity_ids.filtered(lambda a: a.activity_type_id == self.activity_type)

    def test_reminds_overdue_drafts_once(self):
        Sheet = self.env['eh.daily.disbursement']
        Sheet._cron_remind_overdue_drafts()
        Sheet._cron_remind_overdue_drafts()
        reminder = self._reminders(self.overdue)
        self.assertEqual(len(reminder), 1)
        self.assertEqual(reminder.user_id, self.user)
        self.assertFalse(self._reminders(self.current))

    def test_confirm_marks_reminder_done_and_cancel_removes_it(self):
        other = self._create_sheet(
            [('out', self.product_a, 1)], date=self.overdue.date - timedelta(days=1),
            responsible_id=self.user.id)
        self.env['eh.daily.disbursement']._cron_remind_overdue_drafts()
        self.overdue.with_user(self.user).action_confirm()
        self.assertFalse(self._reminders(self.overdue))
        self.assertTrue(self.overdue.message_ids.filtered(
            lambda m: m.mail_activity_type_id == self.activity_type))
        other.with_user(self.user).action_cancel()
        self.assertFalse(self._reminders(other))

    def test_non_draft_sheets_not_reminded(self):
        self.overdue.action_confirm()
        self.env['eh.daily.disbursement']._cron_remind_overdue_drafts()
        self.assertFalse(self._reminders(self.overdue))

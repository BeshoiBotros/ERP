# -*- encoding: utf-8 -*-
##############################################################################
#
# ERP Heritage
# Copyright (C) 2026 (https://www.erpheritage.com.au/)
#
##############################################################################
"""PDF reports (sheet, period summary), XLSX export and its download URL."""

import io
from datetime import date

from openpyxl import load_workbook

from odoo.tests import HttpCase, new_test_user, tagged

from .common import EhDailyDisbursementCase

PERIOD = {'period': 'custom', 'date_from': '2026-09-01', 'date_to': '2026-09-03'}


@tagged('post_install', '-at_install', 'eh_daily_disbursement')
class TestReports(EhDailyDisbursementCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.sheet = cls._create_sheet(
            [('out', cls.product_a, 3), ('out', cls.product_b, 2), ('return', cls.product_a, 1)],
            date=date(2026, 9, 1))                         # +30 +10 -10
        cls.sheet.action_confirm()
        cls._create_sheet([('out', cls.product_b, 4)], date=date(2026, 9, 3))  # +20, draft
        # Previous window (2026-08-29..31): net 20, so the summary shows a change.
        cls._create_sheet([('out', cls.product_a, 2)], date=date(2026, 8, 30)).action_confirm()

    def _html(self, report, docids, data=None):
        html, _fmt = self.env['ir.actions.report']._render_qweb_html(report, docids, data=data)
        return html.decode()

    def test_sheet_pdf_content(self):
        html = self._html('eh_daily_disbursement.action_report_daily_disbursement_sheet', self.sheet.ids)
        for text in (self.sheet.name, 'Daily Disbursement Sheet', self.product_a.display_name,
                     self.product_b.display_name, 'Total disbursed', 'Total returned',
                     'Prepared by', 'Confirmed by'):
            self.assertIn(text, html)

    def test_sheet_report_in_print_menu(self):
        report = self.env.ref('eh_daily_disbursement.action_report_daily_disbursement_sheet')
        self.assertEqual(report.binding_model_id.model, 'eh.daily.disbursement')
        self.assertEqual(report.binding_type, 'report')

    def test_summary_pdf_matches_dashboard(self):
        html = self._html('eh_daily_disbursement.action_report_daily_disbursement_summary', [],
                          data=dict(PERIOD, include_draft=True))
        for text in ('Disbursement Summary', 'All tags', 'Included', 'Disbursed products',
                     'Returned products', self.product_a.display_name, 'By day'):
            self.assertIn(text, html)
        # Net 50 (30 + 10 + 20 - 10) against 20 before: +150%.
        self.assertIn('50.00', html)
        self.assertIn('(+150.0%)', html)
        without_drafts = self._html('eh_daily_disbursement.action_report_daily_disbursement_summary', [],
                                    data=dict(PERIOD, include_draft=False, tag_id=self.tag.id))
        self.assertIn('Excluded', without_drafts)
        self.assertIn(self.tag.name, without_drafts)

    def test_xlsx_export(self):
        filename, content = self.env['eh.daily.disbursement.dashboard']._export_xlsx(
            include_draft=True, **PERIOD)
        self.assertEqual(filename, 'daily_disbursement_2026-09-01_2026-09-03.xlsx')
        workbook = load_workbook(io.BytesIO(content))
        self.assertEqual(workbook.sheetnames, ['Summary', 'By Day', 'Products', 'By Tag', 'Lines'])

        summary = {row[0]: row[1:] for row in workbook['Summary'].iter_rows(values_only=True) if row[0]}
        self.assertEqual(summary['Net cost'][0], 50)
        self.assertEqual(summary['Disbursed cost'][0], 60)

        by_day = list(workbook['By Day'].iter_rows(values_only=True))
        self.assertEqual(len(by_day), 1 + 3 + 1)            # header, every day, total
        self.assertEqual(by_day[1][0].date(), date(2026, 9, 1))   # a real date, sortable in Excel
        self.assertEqual(by_day[1][1:], (40, 10, 30))
        self.assertEqual(by_day[2][1:], (0, 0, 0))          # empty day still listed
        self.assertEqual(by_day[-1][3], '=SUM(D2:D4)')

        products = {row[0]: row for row in workbook['Products'].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(products[self.product_b.display_name][2:4], (6, 30))

        lines = list(workbook['Lines'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(lines), 4)
        self.assertEqual({l[5] for l in lines}, {'Disbursement', 'Return'})
        self.assertEqual(sum(l[10] for l in lines), 50)     # signed amounts add up to net

    def test_xlsx_respects_filters(self):
        _name, content = self.env['eh.daily.disbursement.dashboard']._export_xlsx(
            include_draft=False, **PERIOD)
        lines = list(load_workbook(io.BytesIO(content))['Lines'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(lines), 3)
        self.assertEqual({l[2] for l in lines}, {'Confirmed'})


@tagged('post_install', '-at_install', 'eh_daily_disbursement')
class TestExportRoute(HttpCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        new_test_user(cls.env, login='dd_export_user',
                      groups='eh_daily_disbursement.group_daily_disbursement_user')
        new_test_user(cls.env, login='dd_export_outsider', groups='base.group_user')
        product = cls.env['product.product'].create({'name': 'Export Item', 'standard_price': 7})
        cls.env['eh.daily.disbursement'].create({
            'date': date(2026, 9, 2),
            'line_ids': [(0, 0, {'product_id': product.id, 'qty': 3})],
        })

    def _get(self, **params):
        query = '&'.join('%s=%s' % kv for kv in dict(PERIOD, **params).items())
        return self.url_open('/eh_daily_disbursement/export/xlsx?' + query)

    def test_user_downloads_workbook(self):
        self.authenticate('dd_export_user', 'dd_export_user')
        response = self._get(include_draft='1')
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response.headers['Content-Type'])
        self.assertIn('daily_disbursement_2026-09-01_2026-09-03.xlsx',
                      response.headers['Content-Disposition'])
        workbook = load_workbook(io.BytesIO(response.content))
        lines = list(workbook['Lines'].iter_rows(min_row=2, values_only=True))
        self.assertEqual([(l[4], l[9]) for l in lines], [('Export Item', 21)])

    def test_outsider_refused(self):
        self.authenticate('dd_export_outsider', 'dd_export_outsider')
        response = self._get()
        self.assertEqual(response.status_code, 403)
